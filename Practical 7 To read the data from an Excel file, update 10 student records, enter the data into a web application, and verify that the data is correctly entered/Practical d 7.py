from openpyxl import load_workbook

wb=load_workbook("student.xlsx")

sheet=wb["Student"]

for row in range(2,12):

   roll=sheet.cell(row,1).value
   name=sheet.cell(row,2).value
   course=sheet.cell(row,3).value

   print(roll,name,course)