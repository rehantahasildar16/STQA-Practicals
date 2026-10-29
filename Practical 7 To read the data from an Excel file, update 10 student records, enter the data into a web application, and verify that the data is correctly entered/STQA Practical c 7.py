from openpyxl import load_workbook

# Open Excel file
wb = load_workbook("practical 7.xlsx")

# Select the first worksheet
sheet = wb.active

# Update only the Course column (Column 4)
for row in range(2, 12):
    sheet.cell(row=row, column=4).value = "MCA"

# Save the file
wb.save("practical 7.xlsx")

print("Course column updated successfully.")